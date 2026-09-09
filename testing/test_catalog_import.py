# -*- coding: utf-8 -*-
"""产品表导入 → 商品库 / 库存 / FOB 阶梯 → 买家真的能问到 的全链路测试"""
import os, sys, io, tempfile
TMP = tempfile.mkdtemp(prefix="dxl_ci_")
os.environ.update(WECOM_CORP_ID="ww", WECOM_KF_SECRET="s", WECOM_TOKEN="t", WECOM_AES_KEY="A"*43,
                  BOSS_KEY="k", LLM_ENABLED="0", AUDIT_SCHEDULER="0", DATA_DIR=TMP)
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT); os.chdir(ROOT)
import openpyxl
import catalog_import as ci
import dianxiaoli_core as dx
from wecom_core import InboundMessage
from fastapi.testclient import TestClient
import wecom_service

P = F = 0
def chk(n, cond, e=""):
    global P, F
    P, F = (P+1, F) if cond else (P, F+1)
    print(("PASS  " if cond else "FAIL  ") + n + ("" if cond else "   <- " + str(e)[:160]))

def m(u, t):
    return InboundMessage(channel="w", msg_id="x", conversation_id="c", sender_id=u, account_id="kfA",
                          msg_type="text", text=t, media_bytes=None, media_id="", raw={})
def say(u, t):
    r = dx.brain(m(u, t)); return r.text if r else ""

def xlsx(rows):
    wb = openpyxl.Workbook(); ws = wb.active
    for r in rows: ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()

# ── 1. 中文表头 + 工厂列 ──
data = xlsx([
    ["产品表", None, None],                                  # 标题行，应跳过
    ["货号", "名称", "零售价", "拿货价", "进价", "起批量", "库存", "FOB起订", "FOB@500", "FOB@2000", "交期(天)", "港口"],
    ["V7", "V7 不锈钢阀门 DN50", 300, 200, 150, 10, 300, 100, 40, 36, 30, "Shanghai"],
    ["T2", "T2 保温杯 350ml", 39, 26, "", 50, 0, "", 4.5, 4.2, "", ""],
    ["", "空行", 1, 1],
    ["T2", "T2 重复行", 39, 25, "", 50, 0],
    ["X9", "没有价格的", "", "", "", "", 5],
])
items, errors = ci.parse_xlsx(data)
chk("跳过标题行找到表头", len(items) == 2, [i["sku"] for i in items])
v7 = next(i for i in items if i["sku"] == "V7")
chk("工厂 FOB 阶梯解析", v7.get("fob_tiers") == [{"min": 500, "usd": 40.0}, {"min": 2000, "usd": 36.0}], v7.get("fob_tiers"))
chk("FOB 起订/交期/港口", v7.get("moq_fob") == 100 and v7.get("lead_days") == 30 and v7.get("port") == "Shanghai", v7)
chk("重复货号提示且后者覆盖", any("重复" in e for e in errors) and next(i for i in items if i["sku"] == "T2")["trade"] == 25, errors)
chk("无价格行跳过并提示", any("X9" in e for e in errors) and all(i["sku"] != "X9" for i in items), errors)

# ── 2. 英文表头 + 缺拿货价自动推 ──
data2 = xlsx([["SKU", "Description", "Retail", "Stock"], ["Q1", "Q1 widget", 100, 50]])
it2, _ = ci.parse_xlsx(data2)
chk("英文表头识别", it2 and it2[0]["sku"] == "Q1" and it2[0]["name"] == "Q1 widget", it2)
chk("缺拿货价按七折推", it2[0]["trade"] == 70.0, it2)

# ── 3. 写进商品库 → 买家能问到 ──
d = dx.load(); applied = ci.apply_to_store(items + it2, d, replace=True); dx.save(d)
chk("apply 计数", applied == {"added": 3, "updated": 0, "total": 3}, applied)
t = say("b1", "V7阀门什么价")
chk("导入后买家问价可答", "300" in t and "200" in t, t)
chk("库存写入", dx.load()["stock"].get("V7") == 300)
t = say("b2", "T2保温杯来10个")
chk("库存 0 的款按断货登记", "断货" in t and "登记" in t, t)

# 再导一次同货号 → 更新不重复
d = dx.load(); applied2 = ci.apply_to_store([{"sku": "V7", "name": "V7 改名", "retail": 310, "trade": 210, "cost": 150, "stock": 280}], d); dx.save(d)
chk("重复导入走更新", applied2["updated"] == 1 and applied2["total"] == 3 and dx.load()["stock"]["V7"] == 280, applied2)

# ── 4. 工厂模式下 FOB 阶梯真的被用上 ──
d = dx.load(); d["factory"] = {"enabled": True}; dx.save(d)
t = say("b3", "FOB price for 2000 pcs V7 valve")
chk("FOB 用表里的阶梯价而不是推算", "36" in t and "2,000" in t and "Shanghai" not in t, t)   # 港口按询盘/默认，不按表
t = say("b3", "FOB for 50 pcs V7")
chk("FOB 起订量取自表", "MOQ" in t and "100" in t, t)

# ── 5. 端点 + 模板 ──
c = TestClient(wecom_service.app)
r = c.post("/boss/import?key=k&replace=1", files={"file": ("p.xlsx", data, "application/octet-stream")})
chk("/boss/import 导入", r.status_code == 200 and r.json()["total"] == 2 and "导入完成" in r.json()["text"], r.text[:200])
chk("/boss/import 鉴权", c.post("/boss/import", files={"file": ("p.xlsx", data)}).status_code == 401)
r = c.get("/boss/import/template?key=k")
chk("模板下载", r.status_code == 200 and len(r.content) > 3000)
wb = openpyxl.load_workbook(io.BytesIO(r.content)); ws = wb["产品表"]
heads = [c.value for c in ws[1]]
chk("模板表头含工厂列", "FOB@2000" in heads and "货号" in heads, heads)
tpl_items, _ = ci.parse_xlsx(r.content)
chk("模板示例行自己能被导入", len(tpl_items) == 2 and tpl_items[0].get("fob_tiers"), tpl_items)

# ── 6. 老板端「产品表」上传区 ──
r = c.get("/boss?key=k")
h = r.text
chk("老板端有产品表上传区", r.status_code == 200 and 'id="importFile"' in h and 'id="importResult"' in h
    and 'id="importBtn"' in h and "/boss/import/template" in h, r.status_code)
chk("老板端无 key 401", c.get("/boss").status_code == 401)
chk("模板链接带 key 且分档口/工厂版", "key='+encodeURIComponent(KEY)" in h
    and "'&factory=0'" in h and "'&factory=1'" in h, "模板链接未由 JS 拼 key")
chk("上传走 multipart 且支持清空重导", "FormData()" in h and "'&replace=1'" in h
    and "/boss/import?key='+encodeURIComponent(KEY)" in h, "上传逻辑缺失")
chk("UI 展示的是端点返回的 text 字段", "j.text" in h and "text" in c.post(
    "/boss/import?key=k", files={"file": ("p.xlsx", data)}).json(), "text 字段对不上")

print(f"\n结果: {P} passed, {F} failed")
sys.exit(1 if F else 0)
